"""Build-only synthetic inputs. No private documents or external services are used."""
import io
import math
import pathlib
import struct
import sys
import tempfile
import wave
import zipfile

from PIL import Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

destination = pathlib.Path(sys.argv[1])
destination.parent.mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory() as temporary:
    root = pathlib.Path(temporary)
    image = Image.new("RGB", (96, 64), (36, 94, 176))
    image.save(root / "사진.png")
    (root / "안내문.txt").write_text("오프라인 변환 시험\n인터넷 연결 없이 실제 파일을 변환합니다.\nOffline conversion test.\n", encoding="utf-8")

    doc = Document()
    doc.add_heading("오프라인 시험 보고서", 0)
    doc.add_paragraph("외부 서버 없이 문서 변환을 확인합니다. Offline test document.")
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("항목", "내용"), ("시험", "한글 표"), ("결과", "확인용")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    doc.save(root / "보고서.docx")

    book = Workbook()
    sheet = book.active
    sheet.title = "시험 명부"
    for row in [("번호", "항목", "수량"), (1, "서류", 3), (2, "자료", 5)]:
        sheet.append(row)
    sheet.column_dimensions["B"].width = 20
    book.create_sheet("둘째 시트").append(["추가", "시험"])
    book.save(root / "업무명부.xlsx")

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "오프라인 업무 안내"
    slide.placeholders[1].text = "첫 실행과 문서 변환 시험\nOffline conversion test"
    slides.save(root / "업무안내.pptx")

    with wave.open(str(root / "녹음.wav"), "wb") as audio:
        audio.setparams((1, 2, 22050, 0, "NONE", "not compressed"))
        audio.writeframes(b"".join(struct.pack("<h", int(6000 * math.sin(2 * math.pi * 440 * t / 22050))) for t in range(22050)))

    with zipfile.ZipFile(root / "묶음.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("test.txt", "offline archive sample")

    with zipfile.ZipFile(root / "안내책.epub", "w") as epub:
        epub.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        epub.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        epub.writestr("content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" unique-identifier="bookid" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="bookid">urn:uuid:offline-test</dc:identifier><dc:title>Offline test</dc:title><dc:language>ko</dc:language><meta property="dcterms:modified">2026-10-19T00:00:00Z</meta></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="chapter"/></spine></package>')
        epub.writestr("chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Offline test</title></head><body><h1>오프라인 시험</h1><p>인터넷 없이 전자책을 변환합니다. Offline conversion test.</p></body></html>')
        epub.writestr("nav.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="chapter.xhtml">Chapter</a></li></ol></nav></body></html>')

    (root / "간단문서.rtf").write_text(r"{\rtf1\ansi Offline RTF document.\par Conversion without an Internet connection.}", encoding="ascii")
    (root / "도형.eps").write_text("%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 100 100\n0.2 0.4 0.8 setrgbcolor\n10 10 80 80 rectfill\nshowpage\n%%EOF\n", encoding="ascii")
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as fixtures:
        for path in sorted(root.iterdir()):
            fixtures.write(path, path.name)
print(f"Synthetic fixtures: {destination.name}, {destination.stat().st_size} bytes")
